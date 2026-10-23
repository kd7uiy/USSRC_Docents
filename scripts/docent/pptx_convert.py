"""PowerPoint deck -> page bundle of slide WebPs with speaker notes.

Slide images come from PowerPoint itself (COM automation) for full fidelity;
notes, hidden flags and embedded videos are read with python-pptx.
"""
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation

from .images import to_webp
from .markup import front_matter, shortcode_text

SLIDE_WIDTH = 1600
VIDEO_EXTS = {".mp4", ".mov", ".m4v", ".wmv", ".avi", ".mpg", ".mpeg"}


@dataclass
class SlideInfo:
    number: int
    notes: str
    hidden: bool
    videos: list = field(default_factory=list)  # media part names, e.g. "media1.mp4"


def read_slides(path: Path) -> list:
    prs = Presentation(str(path))
    slides = []
    for i, slide in enumerate(prs.slides, start=1):
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        videos = []
        for rel in slide.part.rels.values():
            if rel.is_external:
                continue
            name = Path(str(rel.target_part.partname)).name
            if Path(name).suffix.lower() in VIDEO_EXTS and name not in videos:
                videos.append(name)
        slides.append(SlideInfo(i, notes, slide._element.get("show") == "0", videos))
    return slides


def export_slide_pngs(path: Path, out_dir: Path) -> list:
    """Export every slide as PNG via PowerPoint. Returns paths in slide order."""
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    app = win32com.client.DispatchEx("PowerPoint.Application")
    pres = app.Presentations.Open(str(path.resolve()), True, False, False)  # ReadOnly, Untitled, WithWindow
    try:
        height = round(SLIDE_WIDTH * pres.PageSetup.SlideHeight / pres.PageSetup.SlideWidth)
        out = []
        for slide in pres.Slides:
            png = out_dir / f"slide-{slide.SlideIndex:03d}.png"
            slide.Export(str(png), "PNG", SLIDE_WIDTH, height)
            out.append(png)
        return out
    finally:
        pres.Close()
        app.Quit()
        pythoncom.CoUninitialize()


def convert(src: Path, bundle: Path, title: str, weight: int, video_links) -> list:
    """video_links(src, slide_no, media_name) -> (url or None). Returns slide info."""
    slides = read_slides(src)
    (bundle / "slides").mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        pngs = export_slide_pngs(src, Path(tmp))
        for png in pngs:
            to_webp(png, bundle / "slides" / (png.stem + ".webp"), max_px=SLIDE_WIDTH)

    parts = [front_matter({"title": title, "weight": weight, "docentType": "presentation"})]
    for s in slides:
        if s.hidden:
            continue
        attrs = [f'n="{s.number}"', f'img="slides/slide-{s.number:03d}.webp"']
        for media in s.videos:
            url = video_links(src, s.number, media)
            attrs.append(f'video="{url}"' if url else f'videoPending="{media}"')
            break  # one video link per slide keeps the caption readable
        parts.append(f"{{{{< slide {' '.join(attrs)} >}}}}\n{shortcode_text(s.notes)}\n{{{{< /slide >}}}}\n\n")
    (bundle / "index.md").write_text("".join(parts), "utf-8")
    return slides
