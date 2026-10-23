from pptx import Presentation
from pptx.util import Inches

from docent.pptx_convert import read_slides


def test_notes_and_hidden(tmp_path):
    prs = Presentation()
    layout = prs.slide_layouts[5]
    s1 = prs.slides.add_slide(layout)
    s1.shapes.title.text = "Intro"
    s1.notes_slide.notes_text_frame.text = "Say hello\nSecond line"
    s2 = prs.slides.add_slide(layout)
    s2._element.set("show", "0")
    path = tmp_path / "d.pptx"
    prs.save(path)

    slides = read_slides(path)
    assert [s.number for s in slides] == [1, 2]
    assert slides[0].notes == "Say hello\nSecond line"
    assert slides[0].hidden is False and slides[1].hidden is True
    assert slides[1].notes == ""
    assert slides[0].videos == []
